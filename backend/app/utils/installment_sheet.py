import os
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.core.config import settings

SHEETS_DIR = os.path.join(settings.LOCAL_STORAGE_PATH, "installment-sheets")
os.makedirs(SHEETS_DIR, exist_ok=True)

LOGO_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "logo-icon.png")

INK = colors.HexColor("#1B1B1B")
NAVY = colors.HexColor("#183B66")
GREEN = colors.HexColor("#64A844")
GREY = colors.HexColor("#888888")


def generate_installment_sheet_pdf(product_name: str, sample_amount: float, product, rows: list[dict], show_savings: bool = False) -> str:
    """
    A printable projected installment sheet for a loan product at a sample
    principal amount — for handing to a prospective customer to show exactly
    what their repayments would look like before they've actually applied.
    """
    filename = f"installment-sheet-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.pdf"
    file_path = os.path.join(SHEETS_DIR, filename)
    c = canvas.Canvas(file_path, pagesize=A4)
    width, height = A4

    c.setFillColor(NAVY)
    c.rect(0, height - 30 * mm, width, 30 * mm, fill=1, stroke=0)
    if os.path.exists(LOGO_PATH):
        logo_size = 18 * mm
        pad = 2 * mm
        logo_x, logo_y = 16 * mm, height - 25 * mm
        c.setFillColor(colors.white)
        c.roundRect(logo_x - pad, logo_y - pad, logo_size + 2 * pad, logo_size + 2 * pad, 2.5 * mm, fill=1, stroke=0)
        c.drawImage(ImageReader(LOGO_PATH), logo_x, logo_y, width=logo_size, height=logo_size, mask="auto", preserveAspectRatio=True)
        text_x = logo_x + logo_size + 6 * mm
    else:
        text_x = 16 * mm
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 17)
    c.drawString(text_x, height - 14 * mm, "UDHAYAM MFI")
    c.setFont("Helvetica", 10)
    c.drawString(text_x, height - 21 * mm, f"Projected Installment Sheet — {product_name}")

    y = height - 38 * mm
    c.setFillColor(INK)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(16 * mm, y, f"Sample principal amount: Rs. {sample_amount:,.2f}")
    c.setFont("Helvetica", 9)
    c.drawRightString(width - 16 * mm, y, f"Generated {datetime.utcnow().strftime('%d %b %Y')}")
    y -= 6 * mm
    interest_label = product.custom_interest_label if product.interest_type.value == "other" else product.interest_type.value.capitalize()
    c.setFont("Helvetica", 9)
    c.setFillColor(GREY)
    c.drawString(16 * mm, y, f"{product.interest_rate_annual}% p.a. ({interest_label}) · {product.tenure_months} months · {product.repayment_frequency} repayments")
    y -= 10 * mm

    c.setStrokeColor(colors.HexColor("#DDDDDD"))
    c.setFont("Helvetica-Bold", 9)
    c.setFillColor(GREY)
    if show_savings:
        col_x = [16, 42, 75, 108, 138, 168]
        headers = ["#", "DUE DATE", "PRINCIPAL", "INTEREST", "SAVINGS", "TOTAL DUE"]
    else:
        col_x = [16, 45, 85, 125, 160]
        headers = ["#", "DUE DATE", "PRINCIPAL", "INTEREST", "TOTAL DUE"]
    for x, h in zip(col_x, headers):
        c.drawString(x * mm, y, h)
    y -= 3 * mm
    c.line(16 * mm, y, width - 16 * mm, y)
    y -= 6 * mm

    c.setFont("Helvetica", 9)
    total_principal = total_interest = total_savings = total_due = 0.0
    for row in rows:
        if y < 20 * mm:
            c.showPage()
            y = height - 20 * mm
            c.setFont("Helvetica-Bold", 9)
            c.setFillColor(GREY)
            for x, h in zip(col_x, headers):
                c.drawString(x * mm, y, h)
            y -= 3 * mm
            c.line(16 * mm, y, width - 16 * mm, y)
            y -= 6 * mm
            c.setFont("Helvetica", 9)
        c.setFillColor(INK)
        c.drawString(col_x[0] * mm, y, str(row["installment_no"]))
        c.drawString(col_x[1] * mm, y, row["due_date"])
        c.drawString(col_x[2] * mm, y, f"Rs. {row['principal_due']:,.2f}")
        c.drawString(col_x[3] * mm, y, f"Rs. {row['interest_due']:,.2f}")
        if show_savings:
            c.drawString(col_x[4] * mm, y, f"Rs. {row.get('savings_due', 0):,.2f}")
            c.drawString(col_x[5] * mm, y, f"Rs. {row['total_due']:,.2f}")
        else:
            c.drawString(col_x[4] * mm, y, f"Rs. {row['total_due']:,.2f}")
        total_principal += row["principal_due"]
        total_interest += row["interest_due"]
        total_savings += row.get("savings_due", 0)
        total_due += row["total_due"]
        y -= 6 * mm

    y -= 2 * mm
    c.setStrokeColor(colors.HexColor("#DDDDDD"))
    c.line(16 * mm, y, width - 16 * mm, y)
    y -= 7 * mm
    c.setFont("Helvetica-Bold", 9)
    c.setFillColor(NAVY)
    c.drawString(col_x[0] * mm, y, "TOTAL")
    c.drawString(col_x[2] * mm, y, f"Rs. {total_principal:,.2f}")
    c.drawString(col_x[3] * mm, y, f"Rs. {total_interest:,.2f}")
    if show_savings:
        c.drawString(col_x[4] * mm, y, f"Rs. {total_savings:,.2f}")
        c.drawString(col_x[5] * mm, y, f"Rs. {total_due:,.2f}")
    else:
        c.drawString(col_x[4] * mm, y, f"Rs. {total_due:,.2f}")

    c.setFillColor(GREY)
    c.setFont("Helvetica-Oblique", 7)
    c.drawCentredString(width / 2, 12 * mm, "This is a projected schedule for illustration — actual figures depend on the disbursal date and exact amount taken.")

    c.showPage()
    c.save()
    return file_path


def generate_installment_sheet_xlsx(product_name: str, sample_amount: float, product, rows: list[dict], show_savings: bool = False) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Installment Sheet"[:31]

    last_col = "F" if show_savings else "E"
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"Udhayam MFI — Projected Installment Sheet — {product_name}"
    ws["A1"].font = Font(bold=True, size=14)
    interest_label = product.custom_interest_label if product.interest_type.value == "other" else product.interest_type.value.capitalize()
    ws["A2"] = f"Sample principal: Rs. {sample_amount:,.2f} — {product.interest_rate_annual}% p.a. ({interest_label}) — {product.tenure_months} months — {product.repayment_frequency}"
    ws["A2"].font = Font(italic=True, size=10, color="666666")

    header_row = 4
    if show_savings:
        headers = ["#", "Due Date", "Principal (Rs.)", "Interest (Rs.)", "Savings (Rs.)", "Total Due (Rs.)"]
    else:
        headers = ["#", "Due Date", "Principal (Rs.)", "Interest (Rs.)", "Total Due (Rs.)"]
    fill = PatternFill(start_color="183B66", end_color="183B66", fill_type="solid")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = fill
        cell.font = Font(color="64A844", bold=True)
        cell.alignment = Alignment(horizontal="center")

    total_col = 6 if show_savings else 5
    total_principal = total_interest = total_savings = total_due = 0.0
    for i, row in enumerate(rows, start=header_row + 1):
        ws.cell(row=i, column=1, value=row["installment_no"])
        ws.cell(row=i, column=2, value=row["due_date"])
        ws.cell(row=i, column=3, value=row["principal_due"]).number_format = "#,##0.00"
        ws.cell(row=i, column=4, value=row["interest_due"]).number_format = "#,##0.00"
        if show_savings:
            ws.cell(row=i, column=5, value=row.get("savings_due", 0)).number_format = "#,##0.00"
        ws.cell(row=i, column=total_col, value=row["total_due"]).number_format = "#,##0.00"
        total_principal += row["principal_due"]
        total_interest += row["interest_due"]
        total_savings += row.get("savings_due", 0)
        total_due += row["total_due"]

    total_row = header_row + len(rows) + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=round(total_principal, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=3).number_format = "#,##0.00"
    ws.cell(row=total_row, column=4, value=round(total_interest, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=4).number_format = "#,##0.00"
    if show_savings:
        ws.cell(row=total_row, column=5, value=round(total_savings, 2)).font = Font(bold=True)
        ws.cell(row=total_row, column=5).number_format = "#,##0.00"
    ws.cell(row=total_row, column=total_col, value=round(total_due, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=total_col).number_format = "#,##0.00"

    widths = [8, 14, 16, 16, 16, 16] if show_savings else [8, 14, 16, 16, 16]
    for col, w in zip("ABCDEF", widths):
        ws.column_dimensions[col].width = w

    filename = f"installment-sheet-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.xlsx"
    file_path = os.path.join(SHEETS_DIR, filename)
    wb.save(file_path)
    return file_path


def generate_loan_installment_sheet_pdf(loan_number: str, payer_name: str, payer_type: str, branch_name: str, is_projected: bool, rows: list[dict], is_group: bool = False, show_savings: bool = False, processing_fee: float | None = None) -> str:
    """
    The real installment sheet for one specific loan — with the actual loan
    number and customer/group on it, not a generic sample. For a group loan,
    every row is one member's individual share of one installment, not just
    a group total, so it's clear who owes what. If the loan hasn't been
    disbursed yet, this is a projection (clearly labeled as such); once
    disbursed, it's built from the real schedule with real due dates.
    """
    filename = f"loan-sheet-{loan_number}-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.pdf"
    file_path = os.path.join(SHEETS_DIR, filename)
    c = canvas.Canvas(file_path, pagesize=A4)
    width, height = A4

    c.setFillColor(NAVY)
    c.rect(0, height - 30 * mm, width, 30 * mm, fill=1, stroke=0)
    if os.path.exists(LOGO_PATH):
        logo_size = 18 * mm
        pad = 2 * mm
        logo_x, logo_y = 16 * mm, height - 25 * mm
        c.setFillColor(colors.white)
        c.roundRect(logo_x - pad, logo_y - pad, logo_size + 2 * pad, logo_size + 2 * pad, 2.5 * mm, fill=1, stroke=0)
        c.drawImage(ImageReader(LOGO_PATH), logo_x, logo_y, width=logo_size, height=logo_size, mask="auto", preserveAspectRatio=True)
        text_x = logo_x + logo_size + 6 * mm
    else:
        text_x = 16 * mm
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 17)
    c.drawString(text_x, height - 14 * mm, "UDHAYAM MFI")
    c.setFont("Helvetica", 10)
    c.drawString(text_x, height - 21 * mm, f"Installment Sheet — {loan_number}" + (" (Projected)" if is_projected else ""))

    y = height - 38 * mm
    c.setFillColor(INK)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(16 * mm, y, f"{payer_name} ({payer_type})")
    c.setFont("Helvetica", 9)
    c.drawRightString(width - 16 * mm, y, f"Generated {datetime.utcnow().strftime('%d %b %Y')}")
    y -= 6 * mm
    c.setFillColor(GREY)
    c.drawString(16 * mm, y, f"Branch: {branch_name}   ·   Loan #: {loan_number}")
    y -= 6 * mm
    if processing_fee:
        c.setFillColor(colors.HexColor("#8A6200"))
        c.drawString(16 * mm, y, f"Processing fee (deducted from disbursed amount): Rs. {processing_fee:,.2f}")
        y -= 6 * mm
    y -= 4 * mm

    if is_projected:
        c.setFillColor(colors.HexColor("#8A6200"))
        c.setFont("Helvetica-Oblique", 8)
        note = "This group hasn't been disbursed yet — dates and per-member shares below are projected, not final." if is_group else "This loan hasn't been disbursed yet — dates below are projected from today, not final."
        c.drawString(16 * mm, y, note)
        y -= 8 * mm

    show_savings = show_savings and not is_group  # group sheets show only each member's total, which already
                                                    # folds savings in — no separate breakdown column there
    if is_group:
        col_x = [14, 40, 75, 150]
        headers = ["#", "DUE DATE", "MEMBER", "AMOUNT DUE"]
    elif show_savings:
        col_x = [14, 38, 68, 98, 128, 155]
        headers = ["#", "DUE DATE", "PRINCIPAL", "INTEREST", "SAVINGS", "TOTAL DUE"]
    else:
        col_x = [16, 45, 85, 125, 160]
        headers = ["#", "DUE DATE", "PRINCIPAL", "INTEREST", "TOTAL DUE"]

    def draw_table_header(y):
        c.setFont("Helvetica-Bold", 8 if is_group else 9)
        c.setFillColor(GREY)
        for x, h in zip(col_x, headers):
            c.drawString(x * mm, y, h)
        y -= 3 * mm
        c.setStrokeColor(colors.HexColor("#DDDDDD"))
        c.line(14 * mm, y, width - 14 * mm, y)
        return y - 6 * mm

    y = draw_table_header(y)

    c.setFont("Helvetica", 8 if is_group else 9)
    total_principal = total_interest = total_savings = total_due = 0.0
    for row in rows:
        if y < 20 * mm:
            c.showPage()
            y = height - 20 * mm
            y = draw_table_header(y)
            c.setFont("Helvetica", 8 if is_group else 9)
        c.setFillColor(colors.HexColor("#5A6B00") if row.get("is_paid") else INK)
        if is_group:
            c.drawString(col_x[0] * mm, y, str(row["installment_no"]))
            c.drawString(col_x[1] * mm, y, row["due_date"])
            c.drawString(col_x[2] * mm, y, str(row.get("member_name", ""))[:26])
            c.drawString(col_x[3] * mm, y, f"Rs. {row['total_due']:,.2f}" + (" (Paid)" if row.get("is_paid") else ""))
        elif show_savings:
            c.drawString(col_x[0] * mm, y, str(row["installment_no"]))
            c.drawString(col_x[1] * mm, y, row["due_date"])
            c.drawString(col_x[2] * mm, y, f"Rs. {row['principal_due']:,.2f}")
            c.drawString(col_x[3] * mm, y, f"Rs. {row['interest_due']:,.2f}")
            c.drawString(col_x[4] * mm, y, f"Rs. {row.get('savings_due', 0):,.2f}")
            c.drawString(col_x[5] * mm, y, f"Rs. {row['total_due']:,.2f}" + (" (Paid)" if row.get("is_paid") else ""))
        else:
            c.drawString(col_x[0] * mm, y, str(row["installment_no"]))
            c.drawString(col_x[1] * mm, y, row["due_date"])
            c.drawString(col_x[2] * mm, y, f"Rs. {row['principal_due']:,.2f}")
            c.drawString(col_x[3] * mm, y, f"Rs. {row['interest_due']:,.2f}")
            c.drawString(col_x[4] * mm, y, f"Rs. {row['total_due']:,.2f}" + (" (Paid)" if row.get("is_paid") else ""))
        total_principal += row["principal_due"]
        total_interest += row["interest_due"]
        total_savings += row.get("savings_due", 0)
        total_due += row["total_due"]
        y -= 5.5 * mm if is_group else 6 * mm

    y -= 2 * mm
    c.setStrokeColor(colors.HexColor("#DDDDDD"))
    c.line(14 * mm, y, width - 14 * mm, y)
    y -= 7 * mm
    c.setFont("Helvetica-Bold", 9)
    c.setFillColor(NAVY)
    c.drawString(col_x[0] * mm, y, "TOTAL")
    if is_group:
        c.drawString(col_x[3] * mm, y, f"Rs. {total_due:,.2f}")
    elif show_savings:
        c.drawString(col_x[2] * mm, y, f"Rs. {total_principal:,.2f}")
        c.drawString(col_x[4] * mm, y, f"Rs. {total_savings:,.2f}")
        c.drawString(col_x[5] * mm, y, f"Rs. {total_due:,.2f}")
    else:
        c.drawString(col_x[2] * mm, y, f"Rs. {total_principal:,.2f}")
        c.drawString(col_x[4] * mm, y, f"Rs. {total_due:,.2f}")

    c.setFillColor(GREY)
    c.setFont("Helvetica-Oblique", 7)
    c.drawCentredString(width / 2, 12 * mm, "Generated by Udhayam MFI")

    c.showPage()
    c.save()
    return file_path


def generate_loan_installment_sheet_xlsx(loan_number: str, payer_name: str, payer_type: str, branch_name: str, is_projected: bool, rows: list[dict], is_group: bool = False, show_savings: bool = False, processing_fee: float | None = None) -> str:
    show_savings = show_savings and not is_group  # see the PDF generator above for why groups skip this column
    wb = Workbook()
    ws = wb.active
    ws.title = "Installment Sheet"[:31]

    ws.merge_cells("A1:F1")
    ws["A1"] = f"Udhayam MFI — Installment Sheet — {loan_number}" + (" (Projected)" if is_projected else "")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{payer_name} ({payer_type}) — Branch: {branch_name}"
    ws["A2"].font = Font(italic=True, size=10, color="666666")
    header_row = 4
    if processing_fee:
        ws["A3"] = f"Processing fee (deducted from disbursed amount): Rs. {processing_fee:,.2f}"
        ws["A3"].font = Font(italic=True, size=10, color="8A6200")
        header_row = 5

    if is_group:
        headers = ["#", "Due Date", "Member", "Amount Due (Rs.)", "Status"]
    elif show_savings:
        headers = ["#", "Due Date", "Principal (Rs.)", "Interest (Rs.)", "Savings (Rs.)", "Total Due (Rs.)", "Status"]
    else:
        headers = ["#", "Due Date", "Principal (Rs.)", "Interest (Rs.)", "Total Due (Rs.)", "Status"]
    fill = PatternFill(start_color="183B66", end_color="183B66", fill_type="solid")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = fill
        cell.font = Font(color="64A844", bold=True)
        cell.alignment = Alignment(horizontal="center")

    total_col = 4 if is_group else (6 if show_savings else 5)
    total_principal = total_savings = total_due = 0.0
    for i, row in enumerate(rows, start=header_row + 1):
        col = 1
        ws.cell(row=i, column=col, value=row["installment_no"]); col += 1
        ws.cell(row=i, column=col, value=row["due_date"]); col += 1
        if is_group:
            ws.cell(row=i, column=col, value=row.get("member_name", "")); col += 1
            ws.cell(row=i, column=col, value=row["total_due"]).number_format = "#,##0.00"; col += 1
        else:
            ws.cell(row=i, column=col, value=row["principal_due"]).number_format = "#,##0.00"; col += 1
            ws.cell(row=i, column=col, value=row["interest_due"]).number_format = "#,##0.00"; col += 1
            if show_savings:
                ws.cell(row=i, column=col, value=row.get("savings_due", 0)).number_format = "#,##0.00"; col += 1
            ws.cell(row=i, column=col, value=row["total_due"]).number_format = "#,##0.00"; col += 1
        ws.cell(row=i, column=col, value="Paid" if row.get("is_paid") else "Unpaid")
        total_principal += row["principal_due"]
        total_savings += row.get("savings_due", 0)
        total_due += row["total_due"]

    total_row = header_row + len(rows) + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    if not is_group:
        ws.cell(row=total_row, column=3, value=round(total_principal, 2)).font = Font(bold=True)
        ws.cell(row=total_row, column=3).number_format = "#,##0.00"
        if show_savings:
            ws.cell(row=total_row, column=5, value=round(total_savings, 2)).font = Font(bold=True)
            ws.cell(row=total_row, column=5).number_format = "#,##0.00"
    ws.cell(row=total_row, column=total_col, value=round(total_due, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=total_col).number_format = "#,##0.00"

    if is_group:
        widths = [8, 14, 20, 16, 10]
    elif show_savings:
        widths = [8, 14, 16, 16, 16, 16, 10]
    else:
        widths = [8, 14, 16, 16, 16, 10]
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w

    filename = f"loan-sheet-{loan_number}-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.xlsx"
    file_path = os.path.join(SHEETS_DIR, filename)
    wb.save(file_path)
    return file_path


# ---------------------------------------------------------------------------
# Center/Group aggregate sheet and per-member M.L.L. sheet — additional,
# separate document types matching the client's own real-world paper
# templates (a "Center" roster sheet and an individual "Member Loan Ledger"
# booklet page per member). These are new functions, requested separately by
# the loans router via a `view` query parameter — they don't replace or
# alter generate_loan_installment_sheet_pdf/xlsx above, which keeps producing
# exactly the per-member flat list it always has for any caller that doesn't
# ask for the new views.
# ---------------------------------------------------------------------------

def _draw_letterhead(c, width, height, title_line: str) -> float:
    """Shared navy header bar + logo used by both new sheet types. Returns
    the y-coordinate to start drawing content below the header."""
    c.setFillColor(NAVY)
    c.rect(0, height - 30 * mm, width, 30 * mm, fill=1, stroke=0)
    if os.path.exists(LOGO_PATH):
        logo_size = 18 * mm
        pad = 2 * mm
        logo_x, logo_y = 16 * mm, height - 25 * mm
        c.setFillColor(colors.white)
        c.roundRect(logo_x - pad, logo_y - pad, logo_size + 2 * pad, logo_size + 2 * pad, 2.5 * mm, fill=1, stroke=0)
        c.drawImage(ImageReader(LOGO_PATH), logo_x, logo_y, width=logo_size, height=logo_size, mask="auto", preserveAspectRatio=True)
        text_x = logo_x + logo_size + 6 * mm
    else:
        text_x = 16 * mm
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 17)
    c.drawString(text_x, height - 14 * mm, "UDHAYAM MICRO FINANCE")
    c.setFont("Helvetica", 10)
    c.drawString(text_x, height - 21 * mm, title_line)
    return height - 38 * mm


def _draw_footer(c, width, branch_name: str, branch_phone: str | None = None, branch_address: str | None = None):
    c.setFillColor(GREY)
    c.setFont("Helvetica-Oblique", 7)
    line = f"UDHAYAM MICRO FINANCE — {branch_name}"
    if branch_phone:
        line += f"  ·  {branch_phone}"
    c.drawCentredString(width / 2, 12 * mm, line)
    if branch_address:
        c.drawCentredString(width / 2, 8.5 * mm, branch_address[:110])


def generate_group_center_sheet_pdf(
    loan_number: str, center_name: str, center_place: str | None, branch_name: str,
    is_projected: bool, total_loan_amount: float, members: list[dict], rows: list[dict],
    branch_phone: str | None = None, branch_address: str | None = None,
) -> str:
    """
    The "Center" sheet — one aggregate view of the whole group's loan, used
    by the field officer for the group meeting. Unlike the per-member flat
    list, every row here is the loan's own EMI/Principal/Interest/Saving
    figures exactly as scheduled — no per-member split — plus a roster of
    every member's name and mobile number at the top.
    """
    filename = f"center-sheet-{loan_number}-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.pdf"
    file_path = os.path.join(SHEETS_DIR, filename)
    c = canvas.Canvas(file_path, pagesize=A4)
    width, height = A4

    y = _draw_letterhead(c, width, height, f"Center Sheet — {loan_number}" + (" (Projected)" if is_projected else ""))

    c.setFillColor(INK)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(16 * mm, y, f"CENTER NAME: {center_name}")
    c.setFont("Helvetica", 9)
    c.drawRightString(width - 16 * mm, y, f"Generated {datetime.utcnow().strftime('%d %b %Y')}")
    y -= 6 * mm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(16 * mm, y, f"CENTER PLACE: {center_place or '—'}")
    y -= 6 * mm
    c.setFont("Helvetica", 9)
    c.setFillColor(GREY)
    c.drawString(16 * mm, y, f"Branch: {branch_name}   ·   Loan #: {loan_number}")
    y -= 8 * mm

    # Member roster
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(16 * mm, y, "MEMBER NAME")
    c.drawString(110 * mm, y, "MOBILE NUMBER")
    y -= 4 * mm
    c.setStrokeColor(colors.HexColor("#DDDDDD"))
    c.line(16 * mm, y, width - 16 * mm, y)
    y -= 5 * mm
    c.setFont("Helvetica", 9)
    c.setFillColor(INK)
    for m in members:
        if y < 20 * mm:
            c.showPage()
            y = height - 20 * mm
        c.drawString(16 * mm, y, str(m.get("name", ""))[:40])
        c.drawString(110 * mm, y, str(m.get("phone") or "—"))
        y -= 5 * mm
    y -= 3 * mm

    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(16 * mm, y, f"TOTAL LOAN AMOUNT: Rs. {total_loan_amount:,.2f}")
    y -= 8 * mm

    if is_projected:
        c.setFillColor(colors.HexColor("#8A6200"))
        c.setFont("Helvetica-Oblique", 8)
        c.drawString(16 * mm, y, "This group hasn't been disbursed yet — dates below are projected, not final.")
        y -= 8 * mm

    col_x = [14, 34, 62, 88, 114, 138, 162, 182]
    headers = ["SL.NO", "DATE", "EMI", "PRINCIPAL", "INTEREST", "SAVING", "TOTAL", "CRO SIGN."]

    def draw_table_header(y):
        c.setFont("Helvetica-Bold", 8)
        c.setFillColor(GREY)
        for x, h in zip(col_x, headers):
            c.drawString(x * mm, y, h)
        y -= 3 * mm
        c.setStrokeColor(colors.HexColor("#DDDDDD"))
        c.line(14 * mm, y, width - 14 * mm, y)
        return y - 6 * mm

    y = draw_table_header(y)
    c.setFont("Helvetica", 8)
    total_emi = total_principal = total_interest = total_saving = total_total = 0.0
    for row in rows:
        if y < 20 * mm:
            c.showPage()
            y = height - 20 * mm
            y = draw_table_header(y)
            c.setFont("Helvetica", 8)
        emi = row["principal_due"] + row["interest_due"]
        c.setFillColor(colors.HexColor("#5A6B00") if row.get("is_paid") else INK)
        c.drawString(col_x[0] * mm, y, str(row["installment_no"]))
        c.drawString(col_x[1] * mm, y, row["due_date"])
        c.drawString(col_x[2] * mm, y, f"{emi:,.2f}")
        c.drawString(col_x[3] * mm, y, f"{row['principal_due']:,.2f}")
        c.drawString(col_x[4] * mm, y, f"{row['interest_due']:,.2f}")
        c.drawString(col_x[5] * mm, y, f"{row.get('savings_due', 0):,.2f}")
        c.drawString(col_x[6] * mm, y, f"{row['total_due']:,.2f}")
        total_emi += emi
        total_principal += row["principal_due"]
        total_interest += row["interest_due"]
        total_saving += row.get("savings_due", 0)
        total_total += row["total_due"]
        y -= 5.5 * mm

    y -= 2 * mm
    c.setStrokeColor(colors.HexColor("#DDDDDD"))
    c.line(14 * mm, y, width - 14 * mm, y)
    y -= 7 * mm
    c.setFont("Helvetica-Bold", 9)
    c.setFillColor(NAVY)
    c.drawString(col_x[0] * mm, y, "TOTAL")
    c.drawString(col_x[2] * mm, y, f"{total_emi:,.2f}")
    c.drawString(col_x[3] * mm, y, f"{total_principal:,.2f}")
    c.drawString(col_x[4] * mm, y, f"{total_interest:,.2f}")
    c.drawString(col_x[5] * mm, y, f"{total_saving:,.2f}")
    c.drawString(col_x[6] * mm, y, f"{total_total:,.2f}")

    _draw_footer(c, width, branch_name, branch_phone, branch_address)
    c.showPage()
    c.save()
    return file_path


def generate_group_center_sheet_xlsx(
    loan_number: str, center_name: str, center_place: str | None, branch_name: str,
    is_projected: bool, total_loan_amount: float, members: list[dict], rows: list[dict],
    branch_phone: str | None = None, branch_address: str | None = None,
) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Center Sheet"[:31]

    ws.merge_cells("A1:H1")
    ws["A1"] = f"UDHAYAM MICRO FINANCE — Center Sheet — {loan_number}" + (" (Projected)" if is_projected else "")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"CENTER NAME: {center_name}    CENTER PLACE: {center_place or '—'}    Branch: {branch_name}"
    ws["A2"].font = Font(italic=True, size=10, color="666666")
    ws["A3"] = f"TOTAL LOAN AMOUNT: Rs. {total_loan_amount:,.2f}"
    ws["A3"].font = Font(bold=True, size=11, color="183B66")

    roster_header_row = 5
    ws.cell(row=roster_header_row, column=1, value="MEMBER NAME").font = Font(bold=True)
    ws.cell(row=roster_header_row, column=2, value="MOBILE NUMBER").font = Font(bold=True)
    r = roster_header_row + 1
    for m in members:
        ws.cell(row=r, column=1, value=m.get("name", ""))
        ws.cell(row=r, column=2, value=m.get("phone") or "—")
        r += 1

    header_row = r + 1
    headers = ["SL.NO", "DATE", "EMI (Rs.)", "PRINCIPAL (Rs.)", "INTEREST (Rs.)", "SAVING (Rs.)", "TOTAL (Rs.)", "CRO SIGNATURE"]
    fill = PatternFill(start_color="183B66", end_color="183B66", fill_type="solid")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = fill
        cell.font = Font(color="64A844", bold=True)
        cell.alignment = Alignment(horizontal="center")

    total_emi = total_principal = total_interest = total_saving = total_total = 0.0
    i = header_row + 1
    for row in rows:
        emi = row["principal_due"] + row["interest_due"]
        ws.cell(row=i, column=1, value=row["installment_no"])
        ws.cell(row=i, column=2, value=row["due_date"])
        ws.cell(row=i, column=3, value=round(emi, 2)).number_format = "#,##0.00"
        ws.cell(row=i, column=4, value=row["principal_due"]).number_format = "#,##0.00"
        ws.cell(row=i, column=5, value=row["interest_due"]).number_format = "#,##0.00"
        ws.cell(row=i, column=6, value=row.get("savings_due", 0)).number_format = "#,##0.00"
        ws.cell(row=i, column=7, value=row["total_due"]).number_format = "#,##0.00"
        total_emi += emi
        total_principal += row["principal_due"]
        total_interest += row["interest_due"]
        total_saving += row.get("savings_due", 0)
        total_total += row["total_due"]
        i += 1

    total_row = i
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=round(total_emi, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=round(total_principal, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=round(total_interest, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=round(total_saving, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=7, value=round(total_total, 2)).font = Font(bold=True)
    for col in (3, 4, 5, 6, 7):
        ws.cell(row=total_row, column=col).number_format = "#,##0.00"

    widths = [8, 14, 14, 16, 14, 12, 14, 16]
    for col, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[col].width = w

    filename = f"center-sheet-{loan_number}-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.xlsx"
    file_path = os.path.join(SHEETS_DIR, filename)
    wb.save(file_path)
    return file_path


def generate_member_mll_sheet_pdf(
    mll_no: str, center_name: str, member_name: str, loan_amount: float,
    loan_dis_date: str | None, mobile_number: str | None, branch_name: str,
    is_projected: bool, rows: list[dict],
    branch_phone: str | None = None, branch_address: str | None = None,
) -> str:
    """
    The per-member "M.L.L." (Member Loan Ledger) sheet — an individual
    booklet page for one group member, with a running BALANCE column that
    steps down from their own loan amount to zero as their principal share
    is repaid week by week.
    """
    filename = f"mll-sheet-{mll_no.replace('/', '-')}-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.pdf"
    file_path = os.path.join(SHEETS_DIR, filename)
    c = canvas.Canvas(file_path, pagesize=A4)
    width, height = A4

    y = _draw_letterhead(c, width, height, "Member Loan Ledger (M.L.L.)" + (" (Projected)" if is_projected else ""))

    c.setFillColor(INK)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(16 * mm, y, f"M.L.L. No: {mll_no}")
    c.setFont("Helvetica", 9)
    c.drawRightString(width - 16 * mm, y, f"Generated {datetime.utcnow().strftime('%d %b %Y')}")
    y -= 6 * mm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(16 * mm, y, f"CENTER NAME: {center_name}")
    y -= 6 * mm
    c.drawString(16 * mm, y, f"MEMBER NAME: {member_name}")
    y -= 6 * mm
    c.setFont("Helvetica", 9)
    c.drawString(16 * mm, y, f"LOAN AMOUNT: Rs. {loan_amount:,.2f}")
    y -= 6 * mm
    c.drawString(16 * mm, y, f"LOAN DIS DATE: {loan_dis_date or '—'}")
    y -= 6 * mm
    c.drawString(16 * mm, y, f"MOBILE NUMBER: {mobile_number or '—'}")
    y -= 6 * mm
    c.setFillColor(GREY)
    c.drawString(16 * mm, y, f"Branch: {branch_name}")
    y -= 8 * mm

    if is_projected:
        c.setFillColor(colors.HexColor("#8A6200"))
        c.setFont("Helvetica-Oblique", 8)
        c.drawString(16 * mm, y, "This loan hasn't been disbursed yet — dates below are projected, not final.")
        y -= 8 * mm

    col_x = [14, 38, 70, 100, 130, 160]
    headers = ["SL.NO", "DATE", "EMI", "PRINCIPAL", "INTEREST", "BALANCE"]

    def draw_table_header(y):
        c.setFont("Helvetica-Bold", 8)
        c.setFillColor(GREY)
        for x, h in zip(col_x, headers):
            c.drawString(x * mm, y, h)
        y -= 3 * mm
        c.setStrokeColor(colors.HexColor("#DDDDDD"))
        c.line(14 * mm, y, width - 14 * mm, y)
        return y - 6 * mm

    y = draw_table_header(y)
    c.setFont("Helvetica", 8)
    for row in rows:
        if y < 20 * mm:
            c.showPage()
            y = height - 20 * mm
            y = draw_table_header(y)
            c.setFont("Helvetica", 8)
        emi = row["principal_due"] + row["interest_due"]
        c.setFillColor(colors.HexColor("#5A6B00") if row.get("is_paid") else INK)
        c.drawString(col_x[0] * mm, y, str(row["installment_no"]))
        c.drawString(col_x[1] * mm, y, row["due_date"])
        c.drawString(col_x[2] * mm, y, f"{emi:,.2f}")
        c.drawString(col_x[3] * mm, y, f"{row['principal_due']:,.2f}")
        c.drawString(col_x[4] * mm, y, f"{row['interest_due']:,.2f}")
        balance = row["balance"]
        c.drawString(col_x[5] * mm, y, "NILL" if balance <= 0 else f"{balance:,.2f}")
        y -= 6 * mm

    c.drawString(col_x[5] * mm - 30 * mm, 16 * mm, "CRO SIGNATURE: ____________________")
    _draw_footer(c, width, branch_name, branch_phone, branch_address)
    c.showPage()
    c.save()
    return file_path


def generate_member_mll_sheet_xlsx(
    mll_no: str, center_name: str, member_name: str, loan_amount: float,
    loan_dis_date: str | None, mobile_number: str | None, branch_name: str,
    is_projected: bool, rows: list[dict],
    branch_phone: str | None = None, branch_address: str | None = None,
) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "M.L.L."[:31]

    ws.merge_cells("A1:F1")
    ws["A1"] = f"UDHAYAM MICRO FINANCE — Member Loan Ledger — {mll_no}" + (" (Projected)" if is_projected else "")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"CENTER: {center_name}    MEMBER: {member_name}    Branch: {branch_name}"
    ws["A2"].font = Font(italic=True, size=10, color="666666")
    ws["A3"] = f"LOAN AMOUNT: Rs. {loan_amount:,.2f}    LOAN DIS DATE: {loan_dis_date or '—'}    MOBILE: {mobile_number or '—'}"
    ws["A3"].font = Font(size=10)

    header_row = 5
    headers = ["SL.NO", "DATE", "EMI (Rs.)", "PRINCIPAL (Rs.)", "INTEREST (Rs.)", "BALANCE (Rs.)"]
    fill = PatternFill(start_color="183B66", end_color="183B66", fill_type="solid")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.fill = fill
        cell.font = Font(color="64A844", bold=True)
        cell.alignment = Alignment(horizontal="center")

    i = header_row + 1
    for row in rows:
        emi = row["principal_due"] + row["interest_due"]
        ws.cell(row=i, column=1, value=row["installment_no"])
        ws.cell(row=i, column=2, value=row["due_date"])
        ws.cell(row=i, column=3, value=round(emi, 2)).number_format = "#,##0.00"
        ws.cell(row=i, column=4, value=row["principal_due"]).number_format = "#,##0.00"
        ws.cell(row=i, column=5, value=row["interest_due"]).number_format = "#,##0.00"
        balance = row["balance"]
        ws.cell(row=i, column=6, value=("NILL" if balance <= 0 else round(balance, 2)))
        if balance > 0:
            ws.cell(row=i, column=6).number_format = "#,##0.00"
        i += 1

    widths = [8, 14, 14, 16, 14, 14]
    for col, w in zip("ABCDEF", widths):
        ws.column_dimensions[col].width = w

    filename = f"mll-sheet-{mll_no.replace('/', '-')}-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.xlsx"
    file_path = os.path.join(SHEETS_DIR, filename)
    wb.save(file_path)
    return file_path
