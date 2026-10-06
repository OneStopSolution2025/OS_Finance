"""
Profit & Loss and Balance Sheet, including the partnership (capital) side.

Purely a read-only report built from records that already exist — loans,
payments, the EMI schedule, the money audit log (salaries) and partners.
It writes nothing and changes no existing calculation.

Basis (stated on every output):
  * Cash basis. Income is recognised when the payment is actually collected.
  * Each payment is split into principal / interest / member-savings using the
    proportions of the installment it was paid against. A group member's late
    penalty is taken out first and reported as penalty income.
  * Member savings (the custom-schedule "Savings" column) is money held for the
    customer, so it is a liability — not income.
  * Processing fees are income when the loan is disbursed.
  * Salaries come from the immutable money audit log.
  * Cash & bank is derived from recorded transactions (capital in, withdrawals,
    disbursals, collections, salaries). It is not read from a bank statement.
  * Partner profit share is pro-rata to net capital, because no agreed profit
    sharing percentage is stored for partners.
"""
import os
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from sqlalchemy import func

from app.core.config import settings
from app.models.audit import MoneyAuditLog, MoneyEventType
from app.models.finance import Loan, Payment, EMISchedule, GroupContribution, Partner
from app.utils.tz import ist_today, ist_day_bounds_utc

REPORTS_DIR = os.path.join(settings.LOCAL_STORAGE_PATH, "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)

CENT = Decimal("0.01")
ZERO = Decimal("0")

BASIS_NOTES = [
    "Cash basis: income is counted when the payment is collected.",
    "Each payment is split into principal, interest and member savings in the proportion of the installment it was paid against.",
    "Member savings are held for customers, so they appear as a liability, not as income.",
    "Cash & bank is derived from recorded transactions, not from a bank statement.",
    "Partner profit share is pro-rata to net capital. No agreed profit-sharing percentage is stored.",
]


def _q(v) -> Decimal:
    return Decimal(str(v or 0)).quantize(CENT, rounding=ROUND_HALF_UP)


def _split_payment(amount: Decimal, emi, contribution) -> dict:
    """Split one payment into principal / interest / savings / penalty."""
    penalty = ZERO
    if contribution is not None and contribution.penalty_amount:
        penalty = min(_q(contribution.penalty_amount), amount)
    rest = amount - penalty
    total_due = _q(emi.total_due) if emi is not None else ZERO
    if emi is None or total_due <= 0:
        return {"principal": rest, "interest": ZERO, "savings": ZERO, "penalty": penalty}
    interest = (rest * _q(emi.interest_due) / total_due).quantize(CENT, rounding=ROUND_HALF_UP)
    savings = (rest * _q(emi.savings_due) / total_due).quantize(CENT, rounding=ROUND_HALF_UP)
    principal = rest - interest - savings
    return {"principal": principal, "interest": interest, "savings": savings, "penalty": penalty}


def _f(v: Decimal) -> float:
    return float(v.quantize(CENT, rounding=ROUND_HALF_UP))


def build_financials(db, tenant_id: str, from_date: date | None = None, to_date: date | None = None) -> dict:
    as_of = to_date or ist_today()
    if from_date and from_date > as_of:
        from_date = as_of
    start_utc = ist_day_bounds_utc(from_date)[0] if from_date else None
    end_utc = ist_day_bounds_utc(as_of)[1]  # exclusive upper bound — whole of the as-of day included

    def in_period(ts: datetime | None) -> bool:
        if ts is None:
            return False
        return (start_utc is None or ts >= start_utc) and ts < end_utc

    # ---- collections, split into components ----
    rows = (
        db.query(Payment, EMISchedule, GroupContribution)
        .outerjoin(EMISchedule, EMISchedule.id == Payment.emi_id)
        .outerjoin(GroupContribution, GroupContribution.id == Payment.group_contribution_id)
        .filter(Payment.tenant_id == tenant_id, Payment.paid_at < end_utc)
        .all()
    )
    cum = {"principal": ZERO, "interest": ZERO, "savings": ZERO, "penalty": ZERO, "total": ZERO}
    per = {"principal": ZERO, "interest": ZERO, "savings": ZERO, "penalty": ZERO, "total": ZERO}
    for payment, emi, contribution in rows:
        amt = _q(payment.amount)
        parts = _split_payment(amt, emi, contribution)
        for k, v in parts.items():
            cum[k] += v
        cum["total"] += amt
        if in_period(payment.paid_at):
            for k, v in parts.items():
                per[k] += v
            per["total"] += amt

    # ---- disbursed loans ----
    loans = (
        db.query(Loan)
        .filter(Loan.tenant_id == tenant_id, Loan.disbursed_at.isnot(None), Loan.disbursed_at < end_utc)
        .all()
    )
    cum_principal_out = sum((_q(l.principal_amount) for l in loans), ZERO)
    cum_cash_out = sum((_q(l.disbursed_amount) for l in loans), ZERO)
    cum_fees = sum((_q(l.processing_fee) for l in loans), ZERO)
    per_fees = sum((_q(l.processing_fee) for l in loans if in_period(l.disbursed_at)), ZERO)
    per_disbursed = sum((_q(l.principal_amount) for l in loans if in_period(l.disbursed_at)), ZERO)

    # ---- salaries (immutable audit log) ----
    salary_q = db.query(MoneyAuditLog).filter(
        MoneyAuditLog.tenant_id == tenant_id,
        MoneyAuditLog.event_type == MoneyEventType.salary_paid,
        MoneyAuditLog.created_at < end_utc,
    ).all()
    cum_salary = sum((_q(e.amount) for e in salary_q), ZERO)
    per_salary = sum((_q(e.amount) for e in salary_q if in_period(e.created_at)), ZERO)

    # ---- profit & loss ----
    income_period = per["interest"] + per["penalty"] + per_fees
    income_cum = cum["interest"] + cum["penalty"] + cum_fees
    net_period = income_period - per_salary
    net_cum = income_cum - cum_salary

    # ---- partners ----
    partners = db.query(Partner).filter(Partner.tenant_id == tenant_id, Partner.is_active == True).order_by(Partner.created_at).all()
    partner_rows = []
    for p in partners:
        invested = _q(p.invested_amount) if (p.invested_date is None or p.invested_date <= as_of) else ZERO
        withdrawn = _q(p.withdrawal_amount) if (p.withdrawal_date is None or p.withdrawal_date <= as_of) else ZERO
        partner_rows.append({"id": p.id, "name": p.name, "invested": invested, "withdrawn": withdrawn, "net": invested - withdrawn})
    total_invested = sum((r["invested"] for r in partner_rows), ZERO)
    total_withdrawn = sum((r["withdrawn"] for r in partner_rows), ZERO)
    net_capital = total_invested - total_withdrawn

    out_partners = []
    for r in partner_rows:
        share = (r["net"] / net_capital) if net_capital > 0 else ZERO
        period_share = (net_period * share).quantize(CENT, rounding=ROUND_HALF_UP)
        retained_share = (net_cum * share).quantize(CENT, rounding=ROUND_HALF_UP)
        out_partners.append({
            "id": r["id"], "name": r["name"],
            "invested": _f(r["invested"]), "withdrawn": _f(r["withdrawn"]), "net_capital": _f(r["net"]),
            "share_pct": round(float(share) * 100, 2),
            "profit_share_period": _f(period_share),
            "profit_share_cumulative": _f(retained_share),
            "total_equity": _f(r["net"] + retained_share),
        })

    # ---- balance sheet (as of the end date) ----
    cash = net_capital - cum_cash_out + cum["total"] - cum_salary
    receivable = cum_principal_out - cum["principal"]
    total_assets = cash + receivable
    liabilities = cum["savings"]
    equity = net_capital + net_cum
    total_le = liabilities + equity

    overdue = (
        db.query(func.coalesce(func.sum(EMISchedule.total_due - EMISchedule.amount_paid), 0))
        .join(Loan, Loan.id == EMISchedule.loan_id)
        .filter(Loan.tenant_id == tenant_id, EMISchedule.is_paid == False, EMISchedule.due_date < as_of)
        .scalar()
    )

    return {
        "meta": {
            "from_date": from_date.isoformat() if from_date else None,
            "to_date": as_of.isoformat(),
            "generated_at": datetime.utcnow().isoformat(),
            "basis": BASIS_NOTES,
        },
        "profit_and_loss": {
            "income": {
                "interest": _f(per["interest"]),
                "processing_fees": _f(per_fees),
                "late_penalties": _f(per["penalty"]),
                "total": _f(income_period),
            },
            "expenses": {"salaries": _f(per_salary), "total": _f(per_salary)},
            "net_profit": _f(net_period),
            "memo": {
                "principal_disbursed": _f(per_disbursed),
                "principal_recovered": _f(per["principal"]),
                "member_savings_collected": _f(per["savings"]),
                "total_collected": _f(per["total"]),
            },
        },
        "balance_sheet": {
            "assets": {"cash_and_bank": _f(cash), "loans_receivable": _f(receivable), "total": _f(total_assets)},
            "liabilities": {"member_savings": _f(liabilities), "total": _f(liabilities)},
            "equity": {
                "partner_capital": _f(net_capital),
                "retained_earnings": _f(net_cum),
                "total": _f(equity),
            },
            "total_liabilities_and_equity": _f(total_le),
            "difference": _f(total_assets - total_le),
            "balanced": abs(total_assets - total_le) < Decimal("0.05"),
            "memo": {"overdue_installments_amount": _f(_q(overdue))},
        },
        "partners": out_partners,
        "partner_totals": {
            "invested": _f(total_invested), "withdrawn": _f(total_withdrawn), "net_capital": _f(net_capital),
            "profit_share_period": _f(sum((_q(p["profit_share_period"]) for p in out_partners), ZERO)),
            "profit_share_cumulative": _f(sum((_q(p["profit_share_cumulative"]) for p in out_partners), ZERO)),
        },
    }


# --------------------------------------------------------------------------- XLSX
NAVY_FILL = PatternFill(start_color="183B66", end_color="183B66", fill_type="solid")
HEAD_FONT = Font(color="64A844", bold=True)
MONEY = "#,##0.00"


def generate_financials_xlsx(tenant_name: str, data: dict) -> str:
    wb = Workbook()
    meta, pl, bs = data["meta"], data["profit_and_loss"], data["balance_sheet"]
    period = f"{meta['from_date'] or 'Beginning'} to {meta['to_date']}"

    def sheet_header(ws, title, subtitle):
        ws.merge_cells("A1:D1")
        ws["A1"] = f"{tenant_name} — {title}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, size=10, color="666666")

    def line(ws, r, label, value=None, bold=False, indent=0):
        c = ws.cell(row=r, column=1, value=("    " * indent) + label)
        c.font = Font(bold=bold)
        if value is not None:
            v = ws.cell(row=r, column=2, value=value)
            v.number_format = MONEY
            v.font = Font(bold=bold)

    def head(ws, r, labels):
        for i, t in enumerate(labels, start=1):
            c = ws.cell(row=r, column=i, value=t)
            c.fill = NAVY_FILL
            c.font = HEAD_FONT
            c.alignment = Alignment(horizontal="center")

    # Profit & Loss
    ws = wb.active
    ws.title = "Profit & Loss"
    sheet_header(ws, "Profit & Loss", f"Period: {period}")
    head(ws, 4, ["Particulars", "Amount (Rs.)"])
    r = 5
    line(ws, r, "INCOME", bold=True); r += 1
    line(ws, r, "Interest income", pl["income"]["interest"], indent=1); r += 1
    line(ws, r, "Processing fees", pl["income"]["processing_fees"], indent=1); r += 1
    line(ws, r, "Late penalties", pl["income"]["late_penalties"], indent=1); r += 1
    line(ws, r, "Total income", pl["income"]["total"], bold=True); r += 2
    line(ws, r, "EXPENSES", bold=True); r += 1
    line(ws, r, "Salaries", pl["expenses"]["salaries"], indent=1); r += 1
    line(ws, r, "Total expenses", pl["expenses"]["total"], bold=True); r += 2
    line(ws, r, "NET PROFIT / (LOSS)", pl["net_profit"], bold=True); r += 2
    line(ws, r, "Memo (not part of profit)", bold=True); r += 1
    line(ws, r, "Principal disbursed", pl["memo"]["principal_disbursed"], indent=1); r += 1
    line(ws, r, "Principal recovered", pl["memo"]["principal_recovered"], indent=1); r += 1
    line(ws, r, "Member savings collected", pl["memo"]["member_savings_collected"], indent=1); r += 1
    line(ws, r, "Total collected", pl["memo"]["total_collected"], indent=1); r += 2
    line(ws, r, "Net profit share by partner (pro-rata to net capital)", bold=True); r += 1
    head(ws, r, ["Partner", "Share %", "Profit share (Rs.)"]); r += 1
    for p in data["partners"]:
        ws.cell(row=r, column=1, value=p["name"])
        ws.cell(row=r, column=2, value=p["share_pct"])
        ws.cell(row=r, column=3, value=p["profit_share_period"]).number_format = MONEY
        r += 1
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 20

    # Balance Sheet
    ws = wb.create_sheet("Balance Sheet")
    sheet_header(ws, "Balance Sheet", f"As of {meta['to_date']}")
    head(ws, 4, ["Particulars", "Amount (Rs.)"])
    r = 5
    line(ws, r, "ASSETS", bold=True); r += 1
    line(ws, r, "Cash & bank (derived)", bs["assets"]["cash_and_bank"], indent=1); r += 1
    line(ws, r, "Loans receivable (principal outstanding)", bs["assets"]["loans_receivable"], indent=1); r += 1
    line(ws, r, "Total assets", bs["assets"]["total"], bold=True); r += 2
    line(ws, r, "LIABILITIES", bold=True); r += 1
    line(ws, r, "Member savings held", bs["liabilities"]["member_savings"], indent=1); r += 1
    line(ws, r, "Total liabilities", bs["liabilities"]["total"], bold=True); r += 2
    line(ws, r, "EQUITY (PARTNERSHIP)", bold=True); r += 1
    line(ws, r, "Partner capital (invested − withdrawn)", bs["equity"]["partner_capital"], indent=1); r += 1
    line(ws, r, "Retained earnings (cumulative profit)", bs["equity"]["retained_earnings"], indent=1); r += 1
    line(ws, r, "Total equity", bs["equity"]["total"], bold=True); r += 2
    line(ws, r, "TOTAL LIABILITIES + EQUITY", bs["total_liabilities_and_equity"], bold=True); r += 1
    line(ws, r, "Difference (should be 0.00)", bs["difference"]); r += 2
    line(ws, r, "Memo: overdue installments not yet collected", bs["memo"]["overdue_installments_amount"]); r += 2
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18

    # Partnership
    ws = wb.create_sheet("Partnership")
    sheet_header(ws, "Partnership capital & profit share", f"As of {meta['to_date']}")
    head(ws, 4, ["Partner", "Invested", "Withdrawn", "Net capital", "Share %", "Profit share (period)", "Profit share (cumulative)", "Total equity"])
    r = 5
    for p in data["partners"]:
        vals = [p["name"], p["invested"], p["withdrawn"], p["net_capital"], p["share_pct"], p["profit_share_period"], p["profit_share_cumulative"], p["total_equity"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            if i not in (1, 5):
                c.number_format = MONEY
        r += 1
    t = data["partner_totals"]
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    for i, v in zip((2, 3, 4, 6, 7), (t["invested"], t["withdrawn"], t["net_capital"], t["profit_share_period"], t["profit_share_cumulative"])):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(bold=True)
        c.number_format = MONEY
    r += 2
    for note in meta["basis"]:
        ws.cell(row=r, column=1, value=note).font = Font(italic=True, size=9, color="666666")
        r += 1
    for col, w in zip("ABCDEFGH", [28, 16, 16, 16, 10, 22, 24, 16]):
        ws.column_dimensions[col].width = w

    path = os.path.join(REPORTS_DIR, f"financial-statements-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.xlsx")
    wb.save(path)
    return path


# --------------------------------------------------------------------------- PDF
NAVY = colors.HexColor("#183B66")
INK = colors.HexColor("#1B1B1B")
GREY = colors.HexColor("#888888")
RED = colors.HexColor("#C0392B")


def generate_financials_pdf(tenant_name: str, data: dict) -> str:
    path = os.path.join(REPORTS_DIR, f"financial-statements-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.pdf")
    c = canvas.Canvas(path, pagesize=A4)
    width, height = A4
    meta, pl, bs = data["meta"], data["profit_and_loss"], data["balance_sheet"]
    period = f"{meta['from_date'] or 'Beginning'} to {meta['to_date']}"

    def header(title, sub):
        c.setFillColor(NAVY)
        c.rect(0, height - 28 * mm, width, 28 * mm, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(16 * mm, height - 14 * mm, tenant_name.upper())
        c.setFont("Helvetica", 10)
        c.drawString(16 * mm, height - 21 * mm, f"{title} — {sub}")
        return height - 38 * mm

    state = {"y": 0}

    def row(label, value=None, bold=False, indent=0, color=INK):
        c.setFillColor(color)
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 9.5)
        c.drawString((18 + indent * 6) * mm, state["y"], label)
        if value is not None:
            c.drawRightString(width - 18 * mm, state["y"], f"Rs. {value:,.2f}")
        state["y"] -= 6 * mm

    def rule():
        c.setStrokeColor(colors.HexColor("#DDDDDD"))
        c.line(16 * mm, state["y"] + 3.5 * mm, width - 16 * mm, state["y"] + 3.5 * mm)

    # Page 1 — Profit & Loss
    state["y"] = header("Profit & Loss", period)
    row("INCOME", bold=True)
    row("Interest income", pl["income"]["interest"], indent=1)
    row("Processing fees", pl["income"]["processing_fees"], indent=1)
    row("Late penalties", pl["income"]["late_penalties"], indent=1)
    rule(); row("Total income", pl["income"]["total"], bold=True)
    state["y"] -= 3 * mm
    row("EXPENSES", bold=True)
    row("Salaries", pl["expenses"]["salaries"], indent=1)
    rule(); row("Total expenses", pl["expenses"]["total"], bold=True)
    state["y"] -= 3 * mm
    rule()
    row("NET PROFIT / (LOSS)", pl["net_profit"], bold=True, color=RED if pl["net_profit"] < 0 else NAVY)
    state["y"] -= 4 * mm
    row("Memo (not part of profit)", bold=True, color=GREY)
    row("Principal disbursed", pl["memo"]["principal_disbursed"], indent=1, color=GREY)
    row("Principal recovered", pl["memo"]["principal_recovered"], indent=1, color=GREY)
    row("Member savings collected", pl["memo"]["member_savings_collected"], indent=1, color=GREY)
    row("Total collected", pl["memo"]["total_collected"], indent=1, color=GREY)
    c.showPage()

    # Page 2 — Balance Sheet + Partnership
    state["y"] = header("Balance Sheet", f"as of {meta['to_date']}")
    row("ASSETS", bold=True)
    row("Cash & bank (derived)", bs["assets"]["cash_and_bank"], indent=1)
    row("Loans receivable (principal outstanding)", bs["assets"]["loans_receivable"], indent=1)
    rule(); row("Total assets", bs["assets"]["total"], bold=True)
    state["y"] -= 3 * mm
    row("LIABILITIES", bold=True)
    row("Member savings held", bs["liabilities"]["member_savings"], indent=1)
    rule(); row("Total liabilities", bs["liabilities"]["total"], bold=True)
    state["y"] -= 3 * mm
    row("EQUITY (PARTNERSHIP)", bold=True)
    row("Partner capital (invested - withdrawn)", bs["equity"]["partner_capital"], indent=1)
    row("Retained earnings (cumulative profit)", bs["equity"]["retained_earnings"], indent=1)
    rule(); row("Total equity", bs["equity"]["total"], bold=True)
    state["y"] -= 3 * mm
    rule()
    row("TOTAL LIABILITIES + EQUITY", bs["total_liabilities_and_equity"], bold=True, color=NAVY)
    state["y"] -= 6 * mm

    c.setFont("Helvetica-Bold", 10)
    c.setFillColor(NAVY)
    c.drawString(16 * mm, state["y"], "PARTNERSHIP")
    state["y"] -= 7 * mm
    cols = [16, 62, 90, 118, 138, 166]
    heads = ["PARTNER", "INVESTED", "WITHDRAWN", "NET CAPITAL", "SHARE %", "PROFIT SHARE"]
    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(GREY)
    for x, h in zip(cols, heads):
        c.drawString(x * mm, state["y"], h)
    state["y"] -= 6 * mm
    c.setFont("Helvetica", 9)
    c.setFillColor(INK)
    for p in data["partners"]:
        if state["y"] < 40 * mm:
            c.showPage()
            state["y"] = height - 25 * mm
            c.setFont("Helvetica", 9)
            c.setFillColor(INK)
        vals = [p["name"][:24], f"{p['invested']:,.0f}", f"{p['withdrawn']:,.0f}", f"{p['net_capital']:,.0f}", f"{p['share_pct']}%", f"{p['profit_share_period']:,.2f}"]
        for x, v in zip(cols, vals):
            c.drawString(x * mm, state["y"], v)
        state["y"] -= 5.5 * mm
    if not data["partners"]:
        c.drawString(16 * mm, state["y"], "No partners recorded.")
        state["y"] -= 5.5 * mm

    c.setFont("Helvetica-Oblique", 7)
    c.setFillColor(GREY)
    yy = 24 * mm
    for note in meta["basis"]:
        c.drawString(16 * mm, yy, note[:150])
        yy -= 3.5 * mm
    c.showPage()
    c.save()
    return path
